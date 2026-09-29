import formulas, sys
xl = formulas.ExcelModel().loads(sys.argv[1]).finish()
sol = xl.calculate()
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1])
ws = wb["Balance"]
def val(sheet, ref):
    for k, v in sol.items():
        if k.upper().endswith(f"[{sys.argv[1].split('/')[-1].upper()}]{sheet.upper()}'!{ref}"):
            return v.value[0][0]
    return None
for r in range(4, ws.max_row + 1):
    lab = ws.cell(row=r, column=1).value
    if lab is None:
        continue
    vals = [val("Balance", f"{c}{r}") for c in "CDEFG"]
    if all(v is None for v in vals):
        print(f"{r:3} {lab}")
        continue
    fmt = lambda v: f"{v:12.3f}" if isinstance(v, (int, float)) else f"{str(v):>12}"
    print(f"{r:3} {lab[:46]:46} " + " ".join(fmt(v) for v in vals))
we = wb["Elements"]
print("\nDecisions:")
for r in range(5, 27):
    print(f"{we.cell(row=r,column=2).value:4}", [val("Elements", f"{c}{r}") for c in "OPQRS"], f"cost={val('Elements', f'M{r}'):.2f}")
