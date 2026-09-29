"""Build PARAMANU Reclaimer mass & energy balance workbook (formula-driven)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.utils import get_column_letter as L

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "paramanu_mass_energy_balance.xlsx")

BLUE = Font(color="1F4E9E")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
HDR_FILL = PatternFill("solid", fgColor="DDE6F0")
IN_FILL = PatternFill("solid", fgColor="FFF7D6")
TOT_FILL = PatternFill("solid", fgColor="EEEEEE")
THIN = Border(top=Side(style="thin"))

wb = Workbook()


def hdr(ws, row, values, start=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=start + i, value=v)
        c.font = BOLD
        c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")


def inp(cell, value, fmt=None):
    cell.value = value
    cell.font = BLUE
    cell.fill = IN_FILL
    if fmt:
        cell.number_format = fmt


# ---------------------------------------------------------------- README
ws = wb.active
ws.title = "README"
readme = [
    ("PARAMANU Reclaimer - Mass & Energy Balance Model (draft v1)", TITLE),
    ("", None),
    ("Basis: 1 wet tonne of excavated landfill material. All flows are per tonne.", None),
    ("Five scenarios along the Commercial <-> Nature slider (s = 0 ... 1) are computed side by side.", None),
    ("", None),
    ("HOW TO USE", BOLD),
    ("1. Edit blue cells on a yellow background (Inputs, Composition, Elements, and slider positions on Balance row 4).", None),
    ("2. Everything else is a formula. Results are on the 'Balance' sheet; charts on 'Charts'.", None),
    ("3. Check Balance row 30 (mass closure) stays at 0.", None),
    ("", None),
    ("SLIDER LOGIC (paper Section 5.2)", BOLD),
    ("Element i is recovered when  Value_i + s * Harm_i  >  Cost_i   (Nature mode ADDS harm avoided; it never ignores value)", None),
    ("Cost_i follows the Sherwood relation: Cost = K * c_eff^(-b), where c_eff = feed concentration x pre-sort concentration factor.", None),
    ("Policy overrides: 'Always' (e.g. chlorine captured by off-gas scrubber) and 'Immobilize' (element deliberately left in slag).", None),
    ("Security-sensitive elements and isotopes (nuclear / radiological) are intentionally NOT listed - see paper Section 8.", None),
    ("Slider-dependent process settings (fines vitrified, extra treatment energy, renewable share) interpolate linearly between the two ends.", None),
    ("", None),
    ("IMPORTANT CAVEATS", BOLD),
    ("* All default numbers are ORDER-OF-MAGNITUDE PLACEHOLDERS for a first draft, not measured data.", None),
    ("* Replace them with site sampling data (ICP-MS, XRF) and vendor figures before publishing.", None),
    ("* Harm_i ($/kg) is a monetized environmental damage cost - the paper must justify these values (Section 5.3).", None),
    ("* Plasma gasification energy figures vary widely in practice; see Teesside (2016) for a cautionary case.", None),
    ("* Gas mass balance is simplified: oxidant is lumped into syngas; flue-gas CO2/H2O not tracked.", None),
    ("* Metal prices are editable assumptions (Elements sheet column F) - update to current market prices.", None),
    ("", None),
    ("COLOR KEY", BOLD),
    ("Blue on yellow = input you can change. Black = formula. Grey rows = totals.", None),
]
for i, (text, font) in enumerate(readme, start=1):
    c = ws.cell(row=i, column=1, value=text)
    if font:
        c.font = font
ws.column_dimensions["A"].width = 130

# ---------------------------------------------------------------- Inputs
wi = wb.create_sheet("Inputs")
wi["A1"] = "Inputs & assumptions"
wi["A1"].font = TITLE
general = [
    # row, label, value, unit, note
    (3, "GENERAL", None, None, None),
    (4, "Feed basis (wet)", 1000, "kg", "1 wet tonne"),
    (5, "Moisture fraction of feed", 0.25, "fraction", "Old landfill material, typical 15-35%"),
    (6, "Oxidant (O2/steam) per kg organic", 0.4, "kg/kg", "Lumped into syngas mass"),
    (7, "Sherwood constant K", 0.01, "$/kg at c = 1", "Calibrate to known recovery costs"),
    (8, "Sherwood exponent b", 0.9, "-", "Literature ~0.8-1.0"),
    (9, "Cold gas efficiency (gasifier)", 0.70, "fraction", "Chemical energy in syngas / in feed"),
    (10, "Syngas -> electricity efficiency", 0.40, "fraction", "Gas engine ~35%, fuel cell ~50%"),
    (11, "H2 share of syngas energy", 0.45, "fraction", "Depends on steam, water-gas shift"),
    (12, "H2 lower heating value", 33.3, "kWh/kg", "Physical constant"),
    (13, "Waste-heat recovery fraction", 0.5, "fraction", "Of non-electric syngas energy"),
    (14, "Latent heat of water", 2.26, "MJ/kg", "Physical constant"),
    (15, "Dryer efficiency", 0.8, "fraction", ""),
    (16, "Aggregate recovery from glass/stone", 0.9, "fraction", "Rest to slag"),
    (17, "Process water make-up", 500, "kg/t", "Washing, leaching, scrubbing"),
    (19, "FIXED ENERGY INTENSITIES", None, None, None),
    (20, "Excavation & handling", 15, "kWh/t wet", "Diesel equivalent"),
    (21, "Pre-sort (shred, screen, sensors, magnets)", 30, "kWh/t wet", ""),
    (22, "Plasma torch electricity", 0.4, "kWh/kg combustible", "Highly vendor-dependent"),
    (23, "Vitrification (melting slag)", 0.6, "kWh/kg slag", "Glass melting ~0.4-0.8"),
    (24, "Off-gas cleaning", 20, "kWh/t wet", "Fans, scrubbers, carbon beds"),
    (26, "ECONOMICS", None, None, None),
    (27, "Tipping fee received", 60, "$/t wet", "Paid by city/county"),
    (28, "Base processing cost (labor, capex, O&M)", 90, "$/t wet", ""),
    (29, "Electricity purchase price", 0.10, "$/kWh", ""),
    (30, "Electricity export price", 0.05, "$/kWh", ""),
    (31, "Soil / aggregate / slag sale price", 0.005, "$/kg", "$5 per tonne"),
]
hdr(wi, 2, ["Parameter", "Value", "Unit", "Note / source"])
for r, lab, val, unit, note in general:
    wi.cell(row=r, column=1, value=lab)
    if val is None:
        wi.cell(row=r, column=1).font = BOLD
        continue
    inp(wi.cell(row=r, column=2), val)
    wi.cell(row=r, column=3, value=unit)
    wi.cell(row=r, column=4, value=note)

hdr(wi, 33, ["SLIDER-DEPENDENT PARAMETERS", "Commercial (s=0)", "Nature (s=1)", "Unit", "Note"])
slider_params = [
    (34, "Fines vitrified (rest reused as cover soil)", 0.1, 0.8, "fraction", "Nature: treat contaminated fines"),
    (35, "Hydrometallurgy + water treatment", 30, 80, "kWh/t wet", ""),
    (36, "PFAS / organic toxin destruction", 0, 40, "kWh/t wet", "Cold plasma or supercritical water"),
    (37, "Electrokinetic treatment of fines", 0, 30, "kWh/t wet", ""),
    (38, "Renewable share of purchased electricity", 0.3, 1.0, "fraction", "Nature end = 100% renewable"),
]
for r, lab, lo, hi, unit, note in slider_params:
    wi.cell(row=r, column=1, value=lab)
    inp(wi.cell(row=r, column=2), lo)
    inp(wi.cell(row=r, column=3), hi)
    wi.cell(row=r, column=4, value=unit)
    wi.cell(row=r, column=5, value=note)
for col, w in zip("ABCDE", [44, 16, 18, 20, 40]):
    wi.column_dimensions[col].width = w

# ---------------------------------------------------------------- Composition
wc = wb.create_sheet("Composition")
wc["A1"] = "Feedstock composition (dry basis)"
wc["A1"].font = TITLE
hdr(wc, 3, ["Component", "Dry fraction", "Dry mass (kg/t)", "LHV (MJ/kg dry)",
            "Ash fraction", "Combustible (1/0)", "Route", "Note",
            "Organic to syngas (kg)", "Chemical energy (kWh)", "Combustible mass (kg)"])
comps = [
    ("Fines / soil", 0.50, 0, 0, 0, "Screen -> reuse as cover soil or vitrify", "ELFM studies: 40-60%"),
    ("Plastics", 0.18, 35, 0.05, 1, "Plasma gasification -> syngas", ""),
    ("Paper / wood / textiles", 0.12, 15, 0.10, 1, "Plasma gasification -> syngas", "Mostly degraded in old cells"),
    ("Glass / stone / ceramics", 0.12, 0, 0, 0, "Wash -> aggregate", ""),
    ("Ferrous metals", 0.04, 0, 0, 0, "Magnet -> smelt", ""),
    ("Non-ferrous metals", 0.012, 0, 0, 0, "Eddy current -> smelt/refine", ""),
    ("E-waste", 0.003, 0, 0, 0, "Sensor sort -> precious metal refining", "Main gold/silver/Pd source"),
    ("Batteries", 0.0005, 0, 0, 0, "Sensor sort -> hydrometallurgy", "Li, Co, Ni, Cd"),
    ("Other / hazardous", 0.0245, 0, 0, 0, "Mixed; to slag after treatment", "Balancing fraction"),
]
C0 = 4
for i, (name, frac, lhv, ash, comb, route, note) in enumerate(comps):
    r = C0 + i
    wc.cell(row=r, column=1, value=name)
    inp(wc.cell(row=r, column=2), frac, "0.00%")
    wc.cell(row=r, column=3, value=f"=B{r}*Inputs!$B$4*(1-Inputs!$B$5)").number_format = "0.00"
    inp(wc.cell(row=r, column=4), lhv)
    inp(wc.cell(row=r, column=5), ash)
    inp(wc.cell(row=r, column=6), comb)
    wc.cell(row=r, column=7, value=route)
    wc.cell(row=r, column=8, value=note)
    wc.cell(row=r, column=9, value=f"=C{r}*F{r}*(1-E{r})").number_format = "0.00"
    wc.cell(row=r, column=10, value=f"=C{r}*F{r}*D{r}/3.6").number_format = "0.0"
    wc.cell(row=r, column=11, value=f"=C{r}*F{r}").number_format = "0.00"
CT = C0 + len(comps)  # totals row (13)
C_END = CT - 1
wc.cell(row=CT, column=1, value="TOTAL").font = BOLD
for col in "BCIJK":
    c = wc[f"{col}{CT}"]
    c.value = f"=SUM({col}{C0}:{col}{C_END})"
    c.font = BOLD
    c.fill = TOT_FILL
wc[f"B{CT}"].number_format = "0.00%"
for col in "CIK":
    wc[f"{col}{CT}"].number_format = "0.00"
wc[f"J{CT}"].number_format = "0.0"
wc.cell(row=CT + 1, column=1, value="Check: dry fractions must sum to 100%")
wc.cell(row=CT + 1, column=2, value=f'=IF(ABS(B{CT}-1)<0.0001,"OK","FIX")')
for col, w in zip("ABCDEFGHIJK", [26, 12, 14, 14, 12, 14, 38, 30, 16, 16, 16]):
    wc.column_dimensions[col].width = w
FINES_ROW, GLASS_ROW = C0, C0 + 3

# ---------------------------------------------------------------- Elements
we = wb.create_sheet("Elements")
we["A1"] = "Element inventory, value, harm, Sherwood cost and slider recovery decisions"
we["A1"].font = TITLE
we["A2"] = ("Concentrations are placeholder estimates of the RECOVERABLE form in dry feed. "
            "Pre-sort factor = how much sorting concentrates the element before extraction.")
SLIDER_COLS = ["C", "D", "E", "F", "G"]  # slider positions on Balance row 4
DEC_COLS = ["O", "P", "Q", "R", "S"]
REC_COLS = ["T", "U", "V", "W", "X"]
hdr(we, 4, ["Element", "Symbol", "Conc. in dry feed (ppm)", "Pre-sort conc. factor",
            "Effective conc. (fraction)", "Value ($/kg)", "Harm if left ($/kg)",
            "Recovery efficiency", "Refining energy (kWh/kg)", "Policy",
            "Product form", "Product / element mass ratio", "Sherwood cost ($/kg)",
            "Mass in feed (kg/t)"]
    + [f"Recover? s{i+1}" for i in range(5)] + [f"Recovered kg s{i+1}" for i in range(5)])
elements = [
    # name, sym, ppm, factor, value, harm, eff, kWh/kg, policy, form, ratio
    ("Iron", "Fe", 35000, 20, 0.30, 0, 0.90, 0.6, "Auto", "Steel scrap / pig iron", 1),
    ("Aluminium", "Al", 6000, 50, 1.50, 0, 0.80, 1.0, "Auto", "Al ingot", 1),
    ("Copper", "Cu", 2500, 30, 9.0, 5, 0.85, 1.5, "Auto", "Cathode copper", 1),
    ("Zinc", "Zn", 1500, 5, 2.7, 10, 0.70, 3.0, "Auto", "Zn oxide / metal", 1),
    ("Lead", "Pb", 800, 10, 2.0, 200, 0.75, 1.0, "Auto", "Lead bullion", 1),
    ("Nickel", "Ni", 300, 20, 16, 20, 0.60, 4.0, "Auto", "Ni metal / sulfate", 1),
    ("Tin", "Sn", 200, 20, 30, 1, 0.60, 3.0, "Auto", "Tin metal", 1),
    ("Antimony", "Sb", 30, 5, 40, 50, 0.50, 5.0, "Auto", "Sb2O3", 1.20),
    ("Chromium", "Cr", 100, 2, 10, 500, 0.50, 5.0, "Auto", "Cr2O3 (stable Cr III)", 1.46),
    ("Cobalt", "Co", 15, 300, 30, 50, 0.60, 10, "Auto", "Co metal / hydroxide", 1),
    ("Lithium", "Li", 20, 500, 60, 5, 0.50, 30, "Auto", "Li2CO3", 5.32),
    ("Rare earths (total)", "REE", 30, 200, 50, 1, 0.40, 50, "Auto", "Mixed RE oxides", 1.17),
    ("Silver", "Ag", 10, 100, 1200, 10, 0.70, 20, "Auto", "Silver bullion", 1),
    ("Gold", "Au", 0.5, 300, 110000, 0, 0.70, 100, "Auto", "Gold bullion", 1),
    ("Palladium", "Pd", 0.2, 300, 35000, 0, 0.60, 100, "Auto", "Pd sponge", 1),
    ("Cadmium", "Cd", 5, 10, 3, 5000, 0.80, 10, "Auto", "CdS (stable)", 1.29),
    ("Mercury", "Hg", 2, 50, 30, 50000, 0.90, 50, "Auto", "HgS (cinnabar)", 1.16),
    ("Arsenic", "As", 10, 1, 1, 3000, 0.60, 10, "Auto", "Scorodite FeAsO4.2H2O", 3.08),
    ("Fluorine", "F", 300, 1, 0.5, 20, 0.70, 2.0, "Auto", "CaF2", 2.05),
    ("Chlorine", "Cl", 5000, 1, 0.05, 1, 0.90, 0.5, "Always", "NaCl (off-gas scrubber)", 1.65),
]
E0 = 5
for i, (name, sym, ppm, fac, val, harm, eff, kwh, pol, form, ratio) in enumerate(elements):
    r = E0 + i
    we.cell(row=r, column=1, value=name)
    we.cell(row=r, column=2, value=sym)
    inp(we.cell(row=r, column=3), ppm, "#,##0.0")
    inp(we.cell(row=r, column=4), fac)
    we.cell(row=r, column=5, value=f"=MIN(1,C{r}*D{r}/1000000)").number_format = "0.00E+00"
    inp(we.cell(row=r, column=6), val, "#,##0.00")
    inp(we.cell(row=r, column=7), harm, "#,##0")
    inp(we.cell(row=r, column=8), eff, "0%")
    inp(we.cell(row=r, column=9), kwh)
    inp(we.cell(row=r, column=10), pol)
    we.cell(row=r, column=11, value=form)
    inp(we.cell(row=r, column=12), ratio)
    we.cell(row=r, column=13, value=f"=Inputs!$B$7*E{r}^(-Inputs!$B$8)").number_format = "#,##0.00"
    we.cell(row=r, column=14, value=f"=C{r}/1000000*Composition!$C${CT}").number_format = "0.0000"
    for dc, rc, sc in zip(DEC_COLS, REC_COLS, SLIDER_COLS):
        s = f"Balance!{sc}$4"
        we[f"{dc}{r}"] = (f'=IF($J{r}="Always",1,IF($J{r}="Immobilize",0,'
                          f'IF($F{r}+{s}*$G{r}>$M{r},1,0)))')
        we[f"{rc}{r}"] = f"=$N{r}*$H{r}*{dc}{r}"
        we[f"{rc}{r}"].number_format = "0.0000"
E_END = E0 + len(elements) - 1
T = E_END + 2  # first totals row
tot_rows = {
    "elem": (T, "Recovered elements (kg, elemental)", None),
    "prod": (T + 1, "Recovered products (kg, sold form)", "L"),
    "rev": (T + 2, "Product revenue ($)", "F"),
    "cost": (T + 3, "Recovery cost - Sherwood ($)", "M"),
    "harm": (T + 4, "Environmental harm avoided ($)", "G"),
    "energy": (T + 5, "Refining energy (kWh)", "I"),
}
for key, (r, label, wcol) in tot_rows.items():
    we.cell(row=r, column=1, value=label).font = BOLD
    for rc in REC_COLS:
        rng = f"{rc}{E0}:{rc}{E_END}"
        f = f"=SUM({rng})" if wcol is None else f"=SUMPRODUCT({rng},${wcol}${E0}:${wcol}${E_END})"
        c = we[f"{rc}{r}"]
        c.value = f
        c.font = BOLD
        c.fill = TOT_FILL
        c.number_format = "#,##0.00"
we.cell(row=T, column=14, value="Count recovered:").font = BOLD
for dc in DEC_COLS:
    we[f"{dc}{T}"] = f"=SUM({dc}{E0}:{dc}{E_END})"
    we[f"{dc}{T}"].font = BOLD
widths = [20, 7, 12, 10, 12, 11, 11, 10, 11, 11, 30, 10, 12, 11] + [9] * 5 + [12] * 5
for i, w in enumerate(widths, start=1):
    we.column_dimensions[L(i)].width = w
we.row_dimensions[4].height = 45
we.freeze_panes = "C5"

# ---------------------------------------------------------------- Balance
wbal = wb.create_sheet("Balance", 1)
wbal["A1"] = "PARAMANU Reclaimer - mass & energy balance per wet tonne"
wbal["A1"].font = TITLE
wbal["A2"] = "Columns = positions on the Commercial <-> Nature slider. Edit slider values in row 4."
hdr(wbal, 4, ["Slider position s", "Unit"])
for sc, v in zip(SLIDER_COLS, [0, 0.25, 0.5, 0.75, 1.0]):
    inp(wbal[f"{sc}4"], v, "0.00")
    wbal[f"{sc}4"].font = Font(bold=True, color="1F4E9E")
for sc, v in zip(SLIDER_COLS, ["Commercial", "", "Balanced", "", "Nature"]):
    wbal[f"{sc}5"] = v
    wbal[f"{sc}5"].font = Font(italic=True)
    wbal[f"{sc}5"].alignment = Alignment(horizontal="center")

rows = []  # (row, label, unit, formula_template or None, style)
# formula templates use {c} for this column, {e} for matching Elements column,
# {s} for slider cell


def interp(r):
    return f"=Inputs!$B${r}+{{c}}$4*(Inputs!$C${r}-Inputs!$B${r})"


R = {}  # name -> row
layout = [
    ("sec", "SLIDER-DEPENDENT SETTINGS"),
    ("fv", "Fines vitrified fraction", "fraction", interp(34), "0%"),
    ("hyd", "Hydromet + water treatment intensity", "kWh/t", interp(35), "0.0"),
    ("pfas", "PFAS / toxin destruction intensity", "kWh/t", interp(36), "0.0"),
    ("ek", "Electrokinetic treatment intensity", "kWh/t", interp(37), "0.0"),
    ("ren", "Renewable share of purchased electricity", "fraction", interp(38), "0%"),
    ("gap",),
    ("sec", "MASS BALANCE"),
    ("sub", "Inputs"),
    ("wet", "Wet feed", "kg", "=Inputs!$B$4", "#,##0.0"),
    ("moist", "   memo: moisture", "kg", "=Inputs!$B$4*Inputs!$B$5", "#,##0.0"),
    ("dry", "   memo: dry solids", "kg", f"=Composition!$C${CT}", "#,##0.0"),
    ("ox", "Oxidant (O2 / steam) to gasifier", "kg", f"=Composition!$I${CT}*Inputs!$B$6", "#,##0.0"),
    ("reag", "Reagents bound into stable compounds", "kg",
     f"=Elements!{{e}}{tot_rows['prod'][0]}-Elements!{{e}}{tot_rows['elem'][0]}", "#,##0.00"),
    ("tin", "TOTAL IN", "kg", "={c}{wet}+{c}{ox}+{c}{reag}", "#,##0.0", "tot"),
    ("sub", "Outputs"),
    ("wv", "Water vapour -> condensed & treated", "kg", "={c}{moist}", "#,##0.0"),
    ("syn", "Syngas (organics + oxidant)", "kg", f"=Composition!$I${CT}+{{c}}{{ox}}", "#,##0.0"),
    ("soil", "Clean fines -> reused as cover soil", "kg", f"=Composition!$C${FINES_ROW}*(1-{{c}}{{fv}})", "#,##0.0"),
    ("agg", "Aggregate (washed glass / stone)", "kg", f"=Composition!$C${GLASS_ROW}*Inputs!$B$16", "#,##0.0"),
    ("prod", "Recovered products (metals + stable compounds)", "kg",
     f"=Elements!{{e}}{tot_rows['prod'][0]}", "#,##0.00"),
    ("slag", "Vitrified slag (balancing item)", "kg",
     f"={{c}}{{dry}}-Composition!$I${CT}-{{c}}{{soil}}-{{c}}{{agg}}-Elements!{{e}}{tot_rows['elem'][0]}", "#,##0.0"),
    ("tout", "TOTAL OUT", "kg", "={c}{wv}+{c}{syn}+{c}{soil}+{c}{agg}+{c}{prod}+{c}{slag}", "#,##0.0", "tot"),
    ("close", "Closure check (out - in), must be 0", "kg", "=ROUND({c}{tout}-{c}{tin},6)", "0.000000"),
    ("gap",),
    ("elem", "Recovered elements (elemental mass)", "kg", f"=Elements!{{e}}{tot_rows['elem'][0]}", "#,##0.00"),
    ("cnt", "Number of elements recovered", "count", f"=Elements!{{d}}{T}", "0"),
    ("slagpct", "Slag as share of dry solids", "%", "={c}{slag}/{c}{dry}", "0.0%"),
    ("gap",),
    ("sec", "HYDROGEN & WATER"),
    ("chem", "Syngas chemical energy", "kWh", f"=Composition!$J${CT}*Inputs!$B$9", "#,##0.0"),
    ("h2", "Hydrogen in syngas", "kg", "={c}{chem}*Inputs!$B$11/Inputs!$B$12", "#,##0.00"),
    ("h2o", "Water formed when that H2 is used (x9)", "kg", "={c}{h2}*9", "#,##0.0"),
    ("mk", "Process water make-up", "kg", "=Inputs!$B$17", "#,##0.0"),
    ("dis", "Water to treatment, then discharge (NPDES)", "kg", "={c}{wv}+{c}{h2o}+{c}{mk}", "#,##0.0"),
    ("gap",),
    ("sec", "ELECTRICITY BALANCE"),
    ("sub", "Demand"),
    ("d1", "Excavation & handling", "kWh", "=Inputs!$B$20*Inputs!$B$4/1000", "#,##0.0"),
    ("d2", "Pre-sort", "kWh", "=Inputs!$B$21*Inputs!$B$4/1000", "#,##0.0"),
    ("d3", "Plasma torches (gasification)", "kWh", f"=Composition!$K${CT}*Inputs!$B$22", "#,##0.0"),
    ("d4", "Vitrification of slag", "kWh", "={c}{slag}*Inputs!$B$23", "#,##0.0"),
    ("d5", "Metal refining", "kWh", f"=Elements!{{e}}{tot_rows['energy'][0]}", "#,##0.0"),
    ("d6", "Hydrometallurgy + water treatment", "kWh", "={c}{hyd}*Inputs!$B$4/1000", "#,##0.0"),
    ("d7", "PFAS / toxin destruction", "kWh", "={c}{pfas}*Inputs!$B$4/1000", "#,##0.0"),
    ("d8", "Electrokinetic fines treatment", "kWh", "={c}{ek}*Inputs!$B$4/1000", "#,##0.0"),
    ("d9", "Off-gas cleaning", "kWh", "=Inputs!$B$24*Inputs!$B$4/1000", "#,##0.0"),
    ("dem", "TOTAL DEMAND", "kWh", "=SUM({c}{d1}:{c}{d9})", "#,##0.0", "tot"),
    ("sub", "Supply"),
    ("sup", "Syngas -> electricity (on-site)", "kWh", "={c}{chem}*Inputs!$B$10", "#,##0.0", "tot"),
    ("net", "NET (supply - demand)", "kWh", "={c}{sup}-{c}{dem}", "#,##0.0", "tot"),
    ("self", "Energy self-sufficiency", "%", "=MIN(1,{c}{sup}/{c}{dem})", "0%"),
    ("ext", "External electricity needed", "kWh", "=MAX(0,-{c}{net})", "#,##0.0"),
    ("extr", "   of which renewable", "kWh", "={c}{ext}*{c}{ren}", "#,##0.0"),
    ("exp", "Surplus exported to grid", "kWh", "=MAX(0,{c}{net})", "#,##0.0"),
    ("gap",),
    ("sec", "HEAT BALANCE"),
    ("th1", "Drying demand (evaporate moisture)", "kWh_th", "={c}{moist}*Inputs!$B$14/3.6/Inputs!$B$15", "#,##0.0"),
    ("th2", "Recoverable waste heat", "kWh_th", "={c}{chem}*(1-Inputs!$B$10)*Inputs!$B$13", "#,##0.0"),
    ("th3", "Heat surplus (+) / deficit (-)", "kWh_th", "={c}{th2}-{c}{th1}", "#,##0.0", "tot"),
    ("gap",),
    ("sec", "ECONOMICS"),
    ("i1", "Product revenue", "$", f"=Elements!{{e}}{tot_rows['rev'][0]}", "#,##0.00"),
    ("i2", "Soil / aggregate / slag sales", "$", "=({c}{soil}+{c}{agg}+{c}{slag})*Inputs!$B$31", "#,##0.00"),
    ("i3", "Tipping fee", "$", "=Inputs!$B$27*Inputs!$B$4/1000", "#,##0.00"),
    ("i4", "Electricity export", "$", "={c}{exp}*Inputs!$B$30", "#,##0.00"),
    ("inc", "TOTAL INCOME", "$", "=SUM({c}{i1}:{c}{i4})", "#,##0.00", "tot"),
    ("k1", "Recovery cost (Sherwood)", "$", f"=Elements!{{e}}{tot_rows['cost'][0]}", "#,##0.00"),
    ("k2", "Base processing cost", "$", "=Inputs!$B$28*Inputs!$B$4/1000", "#,##0.00"),
    ("k3", "Electricity purchase", "$", "={c}{ext}*Inputs!$B$29", "#,##0.00"),
    ("cost", "TOTAL COST", "$", "=SUM({c}{k1}:{c}{k3})", "#,##0.00", "tot"),
    ("marg", "NET MARGIN", "$", "={c}{inc}-{c}{cost}", "#,##0.00", "tot"),
    ("harm", "Environmental harm avoided (monetized)", "$", f"=Elements!{{e}}{tot_rows['harm'][0]}", "#,##0.00"),
    ("soc", "Societal value (margin + harm avoided)", "$", "={c}{marg}+{c}{harm}", "#,##0.00", "tot"),
]
r = 6
for item in layout:
    r += 1
    kind = item[0]
    if kind == "gap":
        continue
    if kind == "sec":
        c = wbal.cell(row=r, column=1, value=item[1])
        c.font = Font(bold=True, color="FFFFFF")
        for col in range(1, 8):
            wbal.cell(row=r, column=col).fill = PatternFill("solid", fgColor="1F4E9E")
        continue
    if kind == "sub":
        wbal.cell(row=r, column=1, value=item[1]).font = Font(bold=True, italic=True)
        continue
    R[kind] = r

# second pass: write formulas now that all row numbers are known
r = 6
for item in layout:
    r += 1
    if item[0] in ("gap", "sec", "sub"):
        continue
    key, label, unit, tmpl, fmt = item[:5]
    style = item[5] if len(item) > 5 else None
    wbal.cell(row=r, column=1, value=label)
    wbal.cell(row=r, column=2, value=unit)
    for sc, ec, dc in zip(SLIDER_COLS, REC_COLS, DEC_COLS):
        f = tmpl.replace("{c}", sc).replace("{e}", ec).replace("{d}", dc)
        for k, rowno in R.items():
            f = f.replace("{" + k + "}", str(rowno))
        cell = wbal[f"{sc}{r}"]
        cell.value = f
        cell.number_format = fmt
        if style == "tot":
            cell.font = BOLD
            cell.fill = TOT_FILL
            cell.border = THIN
    if style == "tot":
        wbal.cell(row=r, column=1).font = BOLD
        wbal.cell(row=r, column=1).fill = TOT_FILL
        wbal.cell(row=r, column=2).fill = TOT_FILL
for col, w in zip("ABCDEFG", [48, 9, 14, 14, 14, 14, 14]):
    wbal.column_dimensions[col].width = w
wbal.freeze_panes = "C6"

# ---------------------------------------------------------------- Charts
wch = wb.create_sheet("Charts", 2)
wch["A1"] = "Slider scenarios (x-axis = slider position, 0 = Commercial, 1 = Nature)"
wch["A1"].font = TITLE
cats = Reference(wbal, min_col=3, max_col=7, min_row=4)


def add_chart(title, ytitle, keys, anchor):
    ch = LineChart()
    ch.title = title
    ch.y_axis.title = ytitle
    ch.x_axis.title = "Slider position s"
    ch.height, ch.width = 8, 16
    for k in keys:
        ref = Reference(wbal, min_col=3, max_col=7, min_row=R[k])
        ch.add_data(ref, from_rows=True, titles_from_data=False)
        ch.series[-1].tx = SeriesLabel(v=wbal.cell(row=R[k], column=1).value.strip())
    ch.set_categories(cats)
    wch.add_chart(ch, anchor)


add_chart("Electricity: demand vs. on-site syngas supply", "kWh per wet tonne", ["dem", "sup"], "A3")
add_chart("Economics per wet tonne", "$ per wet tonne", ["marg", "harm", "soc"], "K3")
add_chart("Where the dry solids go", "kg per wet tonne", ["soil", "slag", "prod"], "A21")
add_chart("External electricity needed", "kWh per wet tonne", ["ext", "extr", "exp"], "K21")

wb.save(OUT)
print("saved", OUT)
print("Balance rows:", {k: v for k, v in R.items() if k in ("close", "dem", "sup", "net", "marg")})
