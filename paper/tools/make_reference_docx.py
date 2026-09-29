"""Build tools/reference.docx: a Word template whose styles mirror paramanu.sty.

Starts from pandoc's default reference.docx and rewrites fonts, sizes, spacing,
page setup, header/footer and the custom styles the Lua filter uses.
Run:  python make_reference_docx.py  (needs python-docx)
"""
import copy
import os
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, Mm, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "pandoc-default-reference.docx")
OUT = os.path.join(HERE, "reference.docx")

SERIF = "Latin Modern Roman"
MONO = "Latin Modern Mono"
MATH = "Latin Modern Math"
NAVY = "00008C"          # kinavy  rgb(0,0,0.55)
KIBACK = "F2F2FF"        # kiback  rgb(0.95,0.95,1)
GREEN = RGBColor(0x21, 0x8C, 0x21)   # citegreen rgb(0.13,0.55,0.13)
LINKBLUE = RGBColor(0x00, 0x00, 0xD9)
URLBLUE = RGBColor(0x00, 0x00, 0x99)
RUNNING = "PARAMANU Reclaimer: Atomize-First Recovery from Landfill to the Periodic Table"
TEXTWIDTH_TW = 9026      # 451 pt in twips

doc = Document(SRC)
styles = doc.styles


def set_fonts(rpr_owner, name):
    rpr = rpr_owner.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for k in list(rfonts.attrib):
        del rfonts.attrib[k]
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)


def pstyle(name, base=None, size=11, bold=False, italic=False, align=None, before=0, after=0,
           line=None, first=None, left=None, right=None, hanging=None, keep_next=False,
           color=None, font=SERIF, stype=WD_STYLE_TYPE.PARAGRAPH):
    try:
        st = styles[name]
    except KeyError:
        st = styles.add_style(name, stype)
    if base is not None:
        st.base_style = styles[base]
    f = st.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color if color is not None else RGBColor(0, 0, 0)
    set_fonts(st.element, font)
    if stype == WD_STYLE_TYPE.PARAGRAPH:
        pf = st.paragraph_format
        if align is not None:
            pf.alignment = align
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        if line is not None:
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(line)
        pf.first_line_indent = Pt(first) if first is not None else None
        if hanging is not None:
            pf.first_line_indent = Pt(-hanging)
            pf.left_indent = Pt(hanging)
        if left is not None:
            pf.left_indent = Pt(left)
        if right is not None:
            pf.right_indent = Pt(right)
        pf.keep_with_next = keep_next
        pf.widow_control = True
    return st


def ppr_border(st, sides, color, sz=6, space=4, shade=None):
    ppr = st.element.get_or_add_pPr()
    old = ppr.find(qn("w:pBdr"))
    if old is not None:
        ppr.remove(old)
    bdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        if side in sides:
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), str(space))
            e.set(qn("w:color"), color)
            bdr.append(e)
    ppr.append(bdr)
    if shade:
        s = OxmlElement("w:shd")
        s.set(qn("w:val"), "clear")
        s.set(qn("w:color"), "auto")
        s.set(qn("w:fill"), shade)
        ppr.append(s)


def tabs(st, stops):
    ppr = st.element.get_or_add_pPr()
    t = OxmlElement("w:tabs")
    for kind, pos in stops:
        e = OxmlElement("w:tab")
        e.set(qn("w:val"), kind)
        e.set(qn("w:pos"), str(pos))
        t.append(e)
    ppr.append(t)


# ---- document defaults: Latin Modern everywhere, no theme fonts ------------
rpr_default = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
rf = rpr_default.find(qn("w:rFonts"))
for k in list(rf.attrib):
    del rf.attrib[k]
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), SERIF)

# ---- body text (LaTeX 11pt: 13.6pt baseline, 17pt paragraph indent) ---------
J = WD_ALIGN_PARAGRAPH.JUSTIFY
C = WD_ALIGN_PARAGRAPH.CENTER
pstyle("Normal", size=11, line=13.6, align=J)
pstyle("Body Text", base="Normal", first=17, align=J, line=13.6)
pstyle("First Paragraph", base="Body Text", first=0, align=J, line=13.6)
pstyle("Compact", base="Body Text", first=0, before=2, after=2, align=J, line=13.6)
pstyle("Block Text", base="Body Text", first=0, left=25, right=25)

# ---- headings (article: \Large, \large, \normalsize bold) -------------------
pstyle("Heading 1", base="Normal", size=14.4, bold=True, before=20, after=10, keep_next=True,
       align=WD_ALIGN_PARAGRAPH.LEFT, line=17.3)
pstyle("Heading 2", base="Normal", size=12, bold=True, before=14, after=6, keep_next=True,
       align=WD_ALIGN_PARAGRAPH.LEFT, line=14.5)
pstyle("Heading 3", base="Normal", size=11, bold=True, before=12, after=4, keep_next=True,
       align=WD_ALIGN_PARAGRAPH.LEFT)
for h in ("Heading 1", "Heading 2", "Heading 3"):
    st = styles[h]
    # drop pandoc's theme colour / font on headings
    rpr = st.element.get_or_add_rPr()
    for tag in ("w:color",):
        e = rpr.find(qn(tag))
        if e is not None:
            e.set(qn("w:val"), "000000")
            for k in ("w:themeColor", "w:themeShade", "w:themeTint"):
                if e.get(qn(k)) is not None:
                    del e.attrib[qn(k)]

# ---- title block --------------------------------------------------------------
pstyle("Paper Title", base="Normal", size=17.28, bold=True, align=C, after=12, line=21)
pstyle("Paper Subtitle", base="Normal", size=12, align=C, after=14, line=18)
pstyle("Paper Author", base="Normal", size=12, align=C, after=12, line=14.5)
pstyle("Paper Date", base="Normal", size=12, align=C, after=30, line=14.5)
pstyle("Abstract Heading", base="Normal", size=10, bold=True, align=C, after=4, keep_next=True)
pstyle("Abstract Text", base="Normal", size=10, align=J, first=15, left=25, right=25, line=12)
pstyle("Keywords", base="Abstract Text", size=10, align=J, first=0, left=25, right=25, before=5, line=12)

# ---- Key Insight box ------------------------------------------------------------
kt = pstyle("Key Insight Title", base="Normal", size=11, bold=True, color=RGBColor(255, 255, 255),
            before=10, after=0, left=6, right=6, keep_next=True, align=WD_ALIGN_PARAGRAPH.LEFT, line=15)
ppr_border(kt, ("top", "left", "right"), NAVY, sz=6, space=3, shade=NAVY)
kb = pstyle("Key Insight Body", base="Normal", size=11, align=J, before=0, after=0, left=6, right=6, line=13.6)
ppr_border(kb, ("left", "bottom", "right"), NAVY, sz=6, space=4, shade=KIBACK)
pstyle("Key Insight After", base="Body Text", first=0, before=10)

# ---- equations, algorithm, figures, captions ------------------------------------
eq = pstyle("Equation", base="Normal", align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=6, line=None)
eq.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
tabs(eq, [("center", TEXTWIDTH_TW // 2), ("right", TEXTWIDTH_TW)])
for name, before, after in (("Algorithm", 8, 8), ("Figure", 6, 0), ("Captioned Figure", 6, 0)):
    st = pstyle(name, base="Normal", align=C, before=before, after=after)
    st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    st.paragraph_format.keep_with_next = True
pstyle("Caption", base="Normal", size=11, align=J, before=6, after=10, line=13.6)
pstyle("Image Caption", base="Caption", size=11, align=J, before=6, after=12, line=13.6)
pstyle("Table Caption", base="Caption", size=11, align=C, before=10, after=6, line=13.6, keep_next=True)
pstyle("Bibliography", base="Normal", size=11, align=J, hanging=15, after=9, line=13.6)
pstyle("Closing", base="Normal", size=10, align=C, before=36, left=90, right=90, first=0)
ppr_border(styles["Closing"], ("top",), "000000", sz=4, space=6)
pstyle("Proof", base="Body Text", first=0, before=4, after=4)
pstyle("Theorem", base="Body Text", first=0, before=6, after=6)

# ---- character styles ------------------------------------------------------------
pstyle("Citation Char", size=11, color=GREEN, stype=WD_STYLE_TYPE.CHARACTER)
pstyle("Ref Char", size=11, color=LINKBLUE, stype=WD_STYLE_TYPE.CHARACTER)
hl = styles["Hyperlink"]
hl.font.color.rgb = URLBLUE
hl.font.underline = False
set_fonts(hl.element, MONO)
vc = styles["Verbatim Char"]
set_fonts(vc.element, MONO)
vc.font.size = Pt(11)

# ---- table style: booktabs look (heavy top/bottom rule, light header rule) -------
tbl = styles["Table"].element
tblpr = tbl.find(qn("w:tblPr"))
for old in tblpr.findall(qn("w:tblBorders")):
    tblpr.remove(old)
b = OxmlElement("w:tblBorders")
for side, sz in (("top", 8), ("bottom", 8)):
    e = OxmlElement(f"w:{side}")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
    b.append(e)
for side in ("left", "right", "insideH", "insideV"):
    e = OxmlElement(f"w:{side}")
    e.set(qn("w:val"), "nil")
    b.append(e)
tblpr.append(b)
jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); tblpr.append(jc)
for old in tblpr.findall(qn("w:tblCellMar")):
    tblpr.remove(old)
mar = OxmlElement("w:tblCellMar")
for side in ("left", "right"):
    e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), "80"); e.set(qn("w:type"), "dxa"); mar.append(e)
tblpr.append(mar)
# header-row rule
for old in tbl.findall(qn("w:tblStylePr")):
    tbl.remove(old)
fr = OxmlElement("w:tblStylePr"); fr.set(qn("w:type"), "firstRow")
frp = OxmlElement("w:rPr"); bb = OxmlElement("w:b"); frp.append(bb)
tcpr = OxmlElement("w:tcPr"); tcb = OxmlElement("w:tcBorders")
e = OxmlElement("w:bottom"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "5"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
tcb.append(e); tcpr.append(tcb)
fr.append(frp); fr.append(tcpr)
tbl.append(fr)
tst = styles["Table"]
tst.font.size = Pt(10)
set_fonts(tst.element, SERIF)

# ---- page setup: A4, 1in margins ---------------------------------------------------
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Inches(1))
sec.header_distance = Inches(0.55)
sec.footer_distance = Inches(0.5)

# header: running title + 0.4pt rule
sec.header.is_linked_to_previous = False
hp = sec.header.paragraphs[0]
hp.text = ""
r = hp.add_run(RUNNING)
r.font.size = Pt(10)
set_fonts(r._element, SERIF)
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
ppr = hp._p.get_or_add_pPr()
bdr = OxmlElement("w:pBdr"); e = OxmlElement("w:bottom")
e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "1"); e.set(qn("w:color"), "000000")
bdr.append(e); ppr.append(bdr)


def field(par, instr, color=None):
    def run():
        rr = OxmlElement("w:r")
        rp = OxmlElement("w:rPr")
        rfn = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            rfn.set(qn(a), SERIF)
        rp.append(rfn)
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "22"); rp.append(sz)
        if color:
            c = OxmlElement("w:color"); c.set(qn("w:val"), color); rp.append(c)
        rr.append(rp)
        return rr
    r1 = run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r1.append(fc)
    r2 = run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r2.append(it)
    r3 = run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r3.append(fc)
    r4 = run(); t = OxmlElement("w:t"); t.text = "1"; r4.append(t)
    r5 = run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r5.append(fc)
    for x in (r1, r2, r3, r4, r5):
        par._p.append(x)


sec.footer.is_linked_to_previous = False
fp = sec.footer.paragraphs[0]
fp.text = ""
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
field(fp, "PAGE")
r = fp.add_run(" of ")
r.font.size = Pt(11)
set_fonts(r._element, SERIF)
field(fp, "NUMPAGES", color="0000D9")

# ---- math font -----------------------------------------------------------------
settings = doc.settings.element
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
mathpr = settings.find(f"{{{M}}}mathPr")
if mathpr is None:
    mathpr = OxmlElement("m:mathPr")
    settings.append(mathpr)
for old in mathpr.findall(f"{{{M}}}mathFont"):
    mathpr.remove(old)
mf = OxmlElement("m:mathFont")
mf.set(qn("m:val"), MATH)
mathpr.insert(0, mf)

doc.save(OUT)
print("wrote", OUT)
